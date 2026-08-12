#!/usr/bin/env python3
"""读取U8原始导出，形成可供artifact-tool生成工作簿的标准JSON。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SUMMARY_FIELDS = [
    "期初数量", "期初单价", "期初金额",
    "收入数量", "收入单价", "收入金额",
    "发出数量", "发出单价", "发出金额",
    "结存数量", "结存单价", "结存金额",
]


def number(value: Any) -> float:
    try:
        if value in (None, ""):
            return 0.0
        result = float(value)
        return 0.0 if not math.isfinite(result) else result
    except (TypeError, ValueError):
        return 0.0


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def within_amount_tolerance(value: float, tolerance: float) -> bool:
    """按可审计精度消除二进制浮点尾差后判断金额容差。"""
    return abs(round(value, 10)) <= tolerance


def apply_movement_exclusions(
    record: dict[str, Any],
    included: dict[str, float],
    excluded: dict[str, float],
    excluded_posted: dict[str, float] | None = None,
) -> dict[str, Any]:
    """保留U8月表原值，建立四类移动剔除桥，并记录CAATS筛选后流水。"""
    item = dict(record)
    excluded_posted = excluded if excluded_posted is None else excluded_posted
    original_fields = {
        "income_quantity": "收入数量",
        "income_amount": "收入金额",
        "issue_quantity": "发出数量",
        "issue_amount": "发出金额",
        "end_quantity": "结存数量",
        "end_amount": "结存金额",
    }
    for suffix, field in original_fields.items():
        item[f"u8_original_{suffix}"] = record[field]

    item["excluded_iq"] = excluded.get("iq", 0.0)
    item["excluded_ia"] = excluded.get("ia", 0.0)
    item["excluded_oq"] = excluded.get("oq", 0.0)
    item["excluded_oa"] = excluded.get("oa", 0.0)
    item["excluded_rows"] = excluded.get("rows", 0.0)

    item["excluded_posted_iq"] = excluded_posted.get("iq", 0.0)
    item["excluded_posted_ia"] = excluded_posted.get("ia", 0.0)
    item["excluded_posted_oq"] = excluded_posted.get("oq", 0.0)
    item["excluded_posted_oa"] = excluded_posted.get("oa", 0.0)

    item["u8_filtered_income_quantity"] = record["收入数量"] - item["excluded_posted_iq"]
    item["u8_filtered_income_amount"] = record["收入金额"] - item["excluded_posted_ia"]
    item["u8_filtered_income_price"] = (
        item["u8_filtered_income_amount"] / item["u8_filtered_income_quantity"]
        if abs(item["u8_filtered_income_quantity"]) > 1e-12 else 0.0
    )
    item["u8_filtered_issue_quantity"] = record["发出数量"] - item["excluded_posted_oq"]
    item["u8_filtered_issue_amount"] = record["发出金额"] - item["excluded_posted_oa"]
    item["u8_filtered_issue_price"] = (
        item["u8_filtered_issue_amount"] / item["u8_filtered_issue_quantity"]
        if abs(item["u8_filtered_issue_quantity"]) > 1e-12 else 0.0
    )
    item["u8_filtered_end_quantity"] = (
        record["结存数量"] - item["excluded_posted_iq"] + item["excluded_posted_oq"]
    )
    item["u8_filtered_end_amount"] = (
        record["结存金额"] - item["excluded_posted_ia"] + item["excluded_posted_oa"]
    )
    item["u8_filtered_end_price"] = (
        item["u8_filtered_end_amount"] / item["u8_filtered_end_quantity"]
        if abs(item["u8_filtered_end_quantity"]) > 1e-12 else 0.0
    )

    item["filtered_ledger_iq"] = included.get("iq", 0.0)
    item["filtered_ledger_ia"] = included.get("ia", 0.0)
    item["filtered_ledger_oq"] = included.get("oq", 0.0)
    item["filtered_ledger_oa"] = included.get("oa", 0.0)
    return item


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def required_columns(headers: Iterable[Any], required: Iterable[str], label: str) -> dict[str, int]:
    mapping = {text(value): index for index, value in enumerate(headers) if text(value)}
    missing = [name for name in required if name not in mapping]
    if missing:
        raise ValueError(f"{label}缺少必需字段：{', '.join(missing)}")
    return mapping


def month_from_filename(path: Path) -> str:
    match = re.search(r"(20\d{2})[._-]?(0[1-9]|1[0-2])", path.stem)
    if not match:
        raise ValueError(f"无法从文件名识别期间：{path.name}")
    return f"{match.group(1)}{match.group(2)}"


def convert_xls(source: Path, output_dir: Path, soffice: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    if source.suffix.lower() == ".xlsx":
        return source
    if source.suffix.lower() != ".xls":
        raise ValueError(f"仅支持XLS/XLSX：{source}")
    target = output_dir / f"{source.stem}.xlsx"
    if target.exists():
        target.unlink()
    completed = subprocess.run(
        [str(soffice), "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
        check=False,
        capture_output=True,
        text=True,
    )
    if completed.returncode != 0 or not target.exists():
        raise RuntimeError(
            f"XLS转换失败：{source}\nstdout={completed.stdout}\nstderr={completed.stderr}"
        )
    return target


def read_master(path: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows))
    index = required_columns(
        headers,
        ["仓库编码", "仓库名称", "计价方式", "仓库核算组", "记入成本"],
        "仓库档案",
    )
    master: list[dict[str, Any]] = []
    by_name: dict[str, dict[str, Any]] = {}
    for row in rows:
        if row[index["仓库编码"]] in (None, ""):
            continue
        record = {name: row[column] for name, column in index.items()}
        record["仓库编码"] = text(record["仓库编码"])
        record["仓库名称"] = text(record["仓库名称"])
        record["仓库核算组"] = text(record.get("仓库核算组"))
        if record.get("停用日期") is not None:
            record["停用日期"] = str(record["停用日期"])[:10]
        master.append(record)
        if record["仓库名称"] in by_name:
            raise ValueError(f"仓库档案仓库名称重复：{record['仓库名称']}")
        by_name[record["仓库名称"]] = record
    return master, by_name


def read_summaries(paths: list[Path], periods: list[str]) -> tuple[list[dict[str, Any]], dict[tuple[str, str], dict[str, Any]]]:
    summaries: list[dict[str, Any]] = []
    by_key: dict[tuple[str, str], dict[str, Any]] = {}
    found_periods: list[str] = []
    for path in sorted(paths):
        month = month_from_filename(path)
        if month not in periods:
            continue
        found_periods.append(month)
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = list(next(rows))
        index = required_columns(
            headers,
            ["存货编码", "存货名称", "存货规格", "存货单位", *SUMMARY_FIELDS],
            path.name,
        )
        for row_number, row in enumerate(rows, start=2):
            code = text(row[index["存货编码"]])
            if not code or code == "合计":
                continue
            record: dict[str, Any] = {
                "month": month,
                "code": code,
                "stock_code": text(row[index["存货代码"]]) if "存货代码" in index else "",
                "name": text(row[index["存货名称"]]),
                "spec": text(row[index["存货规格"]]),
                "unit": text(row[index["存货单位"]]),
                "weight": row[index["存货重量"]] if "存货重量" in index else None,
                "reg": text(row[index["存货产品注册证号"]]) if "存货产品注册证号" in index else "",
                "manufacturer": text(row[index["存货生产厂商"]]) if "存货生产厂商" in index else "",
                "source_file": path.name,
                "source_row": row_number,
            }
            for field in SUMMARY_FIELDS:
                record[field] = number(row[index[field]])
            key = (month, code)
            if key in by_key:
                raise ValueError(f"收发存存在重复物料月份：{month}/{code}")
            summaries.append(record)
            by_key[key] = record
    if sorted(found_periods) != sorted(periods):
        raise ValueError(f"收发存期间不完整。期望={periods}，实际={sorted(found_periods)}")
    return summaries, by_key


def build_transition_checks(
    material_rows: list[dict[str, Any]], periods: list[str], quantity_tolerance: float, amount_tolerance: float
) -> list[dict[str, Any]]:
    row_by_key = {(row["month"], row["code"]): row for row in material_rows}
    codes = sorted({row["code"] for row in material_rows})
    checks: list[dict[str, Any]] = []
    for position in range(1, len(periods)):
        previous_month = periods[position - 1]
        month = periods[position]
        quantity_mismatch = 0
        amount_mismatch = 0
        quantity_abs = 0.0
        amount_abs = 0.0
        compared = 0
        for code in codes:
            previous = row_by_key.get((previous_month, code))
            current = row_by_key.get((month, code))
            if previous is None and current is None:
                continue
            compared += 1
            quantity_difference = number(current and current["期初数量"]) - number(previous and previous["结存数量"])
            amount_difference = number(current and current["期初金额"]) - number(previous and previous["结存金额"])
            quantity_abs += abs(quantity_difference)
            amount_abs += abs(amount_difference)
            quantity_mismatch += abs(quantity_difference) > quantity_tolerance
            amount_mismatch += abs(amount_difference) > amount_tolerance
        checks.append(
            {
                "previous_month": previous_month,
                "month": month,
                "compared": compared,
                "quantity_mismatch": quantity_mismatch,
                "amount_mismatch": amount_mismatch,
                "quantity_abs": quantity_abs,
                "amount_abs": amount_abs,
                "status": "PASS" if quantity_mismatch == 0 and amount_mismatch == 0 else "REVIEW",
            }
        )
    return checks


def build_complete_material_panel(
    material_rows: list[dict[str, Any]],
    periods: list[str],
    movement_by_key: dict[tuple[str, str], dict[str, float]] | None = None,
    quantity_tolerance: float = 0.000001,
    amount_tolerance: float = 0.01,
) -> list[dict[str, Any]]:
    """从物料首次出现月份补齐至最终期间；缺月仅在确认无有效流水时按零收发承接。"""
    movement_by_key = movement_by_key or {}
    period_position = {month: index for index, month in enumerate(periods)}
    rows_by_code: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in material_rows:
        rows_by_code[row["code"]][row["month"]] = row

    panel: list[dict[str, Any]] = []
    for code in sorted(rows_by_code):
        real_rows = rows_by_code[code]
        first_position = min(period_position[month] for month in real_rows)
        previous: dict[str, Any] | None = None
        for month in periods[first_position:]:
            real = real_rows.get(month)
            if real is not None:
                current = dict(real)
                current["record_type"] = "原始月度记录"
            else:
                if previous is None:
                    raise ValueError(f"补齐月份缺少前期记录：{month}/{code}")
                movement = movement_by_key.get((month, code), {})
                has_quantity = abs(movement.get("iq", 0.0)) > quantity_tolerance or abs(movement.get("oq", 0.0)) > quantity_tolerance
                has_amount = abs(movement.get("ia", 0.0)) > amount_tolerance or abs(movement.get("oa", 0.0)) > amount_tolerance
                if has_quantity or has_amount:
                    raise ValueError(f"月度汇总缺行但CAATS存在有效收发，不能按零补齐：{month}/{code}")
                current = dict(previous)
                current.update(
                    {
                        "month": month,
                        "source_file": "(补齐无收发月份)",
                        "source_row": None,
                        "record_type": "补齐无收发月份",
                        "期初数量": previous["结存数量"],
                        "期初金额": previous["结存金额"],
                        "期初单价": previous["结存金额"] / previous["结存数量"] if abs(previous["结存数量"]) > 1e-12 else 0.0,
                        "收入数量": 0.0,
                        "收入单价": 0.0,
                        "收入金额": 0.0,
                        "发出数量": 0.0,
                        "发出单价": 0.0,
                        "发出金额": 0.0,
                        "结存数量": previous["结存数量"],
                        "结存金额": previous["结存金额"],
                        "结存单价": previous["结存金额"] / previous["结存数量"] if abs(previous["结存数量"]) > 1e-12 else 0.0,
                        "u8_original_income_quantity": 0.0,
                        "u8_original_income_amount": 0.0,
                        "u8_original_issue_quantity": 0.0,
                        "u8_original_issue_amount": 0.0,
                        "u8_original_end_quantity": previous["结存数量"],
                        "u8_original_end_amount": previous["结存金额"],
                        "excluded_iq": 0.0,
                        "excluded_ia": 0.0,
                        "excluded_oq": 0.0,
                        "excluded_oa": 0.0,
                        "excluded_rows": 0.0,
                        "excluded_posted_iq": 0.0,
                        "excluded_posted_ia": 0.0,
                        "excluded_posted_oq": 0.0,
                        "excluded_posted_oa": 0.0,
                        "u8_filtered_income_quantity": 0.0,
                        "u8_filtered_income_amount": 0.0,
                        "u8_filtered_income_price": 0.0,
                        "u8_filtered_issue_quantity": 0.0,
                        "u8_filtered_issue_amount": 0.0,
                        "u8_filtered_issue_price": 0.0,
                        "u8_filtered_end_quantity": previous["结存数量"],
                        "u8_filtered_end_amount": previous["结存金额"],
                        "u8_filtered_end_price": previous["结存金额"] / previous["结存数量"] if abs(previous["结存数量"]) > 1e-12 else 0.0,
                        "filtered_ledger_iq": 0.0,
                        "filtered_ledger_ia": 0.0,
                        "filtered_ledger_oq": 0.0,
                        "filtered_ledger_oa": 0.0,
                        "caats_avg": previous["结存金额"] / previous["结存数量"] if abs(previous["结存数量"]) > 1e-12 else 0.0,
                        "caats_issue": 0.0,
                        "issue_diff": 0.0,
                        "caats_end": previous["结存金额"],
                        "end_diff": 0.0,
                        "income_amount_gap": 0.0,
                        "issue_amount_gap": 0.0,
                        "reason": "补齐无收发月份；期初期末承接上期",
                        "status": "PASS",
                    }
                )
            panel.append(current)
            previous = current
    return panel


def build_continuous_analysis(
    material_rows: list[dict[str, Any]], periods: list[str], amount_tolerance: float = 0.01
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    row_by_key = {(row["month"], row["code"]): row for row in material_rows}
    codes = sorted({row["code"] for row in material_rows})
    rolled: dict[str, dict[str, Any]] = {}
    result: list[dict[str, Any]] = []
    detail: list[dict[str, Any]] = []
    for month_position, month in enumerate(periods):
        u8_issue = 0.0
        independent_caats_issue = 0.0
        independent_difference = 0.0
        independent_abs = 0.0
        continuous_caats_issue = 0.0
        continuous_difference = 0.0
        continuous_abs = 0.0
        continuous_end_difference = 0.0
        continuous_end_abs = 0.0
        rows = 0
        resets = 0
        for code in codes:
            record = row_by_key.get((month, code))
            if record is None:
                continue
            rows += 1
            prior = rolled.get(code)
            can_roll = prior is not None
            begin_quantity = prior["end_quantity"] if can_roll else record["期初数量"]
            begin_amount = prior["end_amount"] if can_roll else record["期初金额"]
            if month_position and not can_roll:
                resets += 1
            caats_income_quantity = record["filtered_ledger_iq"]
            caats_income_amount = record["filtered_ledger_ia"]
            caats_issue_quantity = record["filtered_ledger_oq"]
            denominator = begin_quantity + caats_income_quantity
            average = (begin_amount + caats_income_amount) / denominator if abs(denominator) > 1e-12 else 0.0
            calculated_issue = caats_issue_quantity * average
            end_quantity = begin_quantity + caats_income_quantity - caats_issue_quantity
            end_amount = begin_amount + caats_income_amount - calculated_issue
            issue_difference = record["u8_filtered_issue_amount"] - calculated_issue
            end_difference = record["u8_filtered_end_amount"] - end_amount
            # 兼容仅用于滚算单元测试/复用的精简记录；正式数据均显式提供U8剔除后收入金额。
            u8_filtered_income_amount = record.get("u8_filtered_income_amount", record.get("filtered_ledger_ia", 0.0))
            income_difference = u8_filtered_income_amount - record["filtered_ledger_ia"]
            begin_difference = record["期初金额"] - begin_amount
            rollforward_check = begin_difference + income_difference - issue_difference - end_difference
            detail.append(
                {
                    "month": month,
                    "code": record["code"],
                    "name": record.get("name", ""),
                    "spec": record.get("spec", ""),
                    "u8_begin_quantity": record["期初数量"],
                    "u8_begin_amount": record["期初金额"],
                    "continuous_begin_quantity": begin_quantity,
                    "continuous_begin_amount": begin_amount,
                    "continuous_begin_difference": begin_difference,
                    "income_amount_difference": income_difference,
                    "u8_filtered_issue_amount": record["u8_filtered_issue_amount"],
                    "continuous_caats_issue": calculated_issue,
                    "continuous_issue_difference": issue_difference,
                    "u8_filtered_end_amount": record["u8_filtered_end_amount"],
                    "continuous_caats_end": end_amount,
                    "continuous_end_difference": end_difference,
                    "rollforward_check": rollforward_check,
                    "rollforward_status": "PASS" if within_amount_tolerance(rollforward_check, amount_tolerance) else "REVIEW",
                    "continuous_average": average,
                    "continuous_end_quantity": end_quantity,
                    "status": "PASS" if within_amount_tolerance(issue_difference, amount_tolerance) else "REVIEW",
                    "start_type": "承接最近可追溯连续期末" if can_roll else "使用U8期初",
                    "record_type": record.get("record_type", "原始月度记录"),
                }
            )
            u8_issue += record["u8_filtered_issue_amount"]
            independent_caats_issue += record["caats_issue"]
            independent_difference += record["issue_diff"]
            independent_abs += abs(record["issue_diff"])
            continuous_caats_issue += calculated_issue
            continuous_difference += issue_difference
            continuous_abs += abs(issue_difference)
            continuous_end_difference += end_difference
            continuous_end_abs += abs(end_difference)
            rolled[code] = {"month": month, "end_quantity": end_quantity, "end_amount": end_amount}
        result.append(
            {
                "month": month,
                "rows": rows,
                "u8_issue": u8_issue,
                "independent_caats_issue": independent_caats_issue,
                "independent_difference": independent_difference,
                "independent_abs": independent_abs,
                "continuous_caats_issue": continuous_caats_issue,
                "continuous_difference": continuous_difference,
                "continuous_abs": continuous_abs,
                "continuous_end_difference": continuous_end_difference,
                "continuous_end_abs": continuous_end_abs,
                "carry_impact": continuous_difference - independent_difference,
                "resets": resets,
            }
        )
    return result, detail


def build_continuous_summary(material_rows: list[dict[str, Any]], periods: list[str]) -> list[dict[str, Any]]:
    result, _ = build_continuous_analysis(material_rows, periods)
    return result


def build_offset_rows(
    material_rows: list[dict[str, Any]], warehouse_rows: list[dict[str, Any]], quantity_tolerance: float, amount_tolerance: float
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    material_by_key = {(row["month"], row["code"]): row for row in material_rows}
    warehouses_by_key: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in warehouse_rows:
        warehouses_by_key[(row["month"], row["code"])].append(row)
    result: list[dict[str, Any]] = []
    multi_issue_keys = 0
    clean_multi_keys = 0
    both_sign_keys = 0
    total_offset = 0.0
    for key, rows in warehouses_by_key.items():
        issue_rows = [row for row in rows if abs(row["oq"]) > quantity_tolerance]
        if len(issue_rows) < 2:
            continue
        multi_issue_keys += 1
        material = material_by_key.get(key)
        if material is None:
            continue
        warehouse_quantity = sum(row["oq"] for row in rows)
        if abs(warehouse_quantity - material["发出数量"]) > quantity_tolerance:
            continue
        if abs(material["issue_amount_gap"]) > amount_tolerance:
            continue
        clean_multi_keys += 1
        parts: list[dict[str, Any]] = []
        for row in rows:
            if abs(row["oq"]) <= quantity_tolerance and abs(row["oa"]) <= amount_tolerance:
                continue
            contribution = row["oa"] - row["oq"] * material["caats_avg"]
            parts.append({"warehouse_row": row, "contribution": contribution})
        positive = sum(max(part["contribution"], 0.0) for part in parts)
        negative_abs = sum(max(-part["contribution"], 0.0) for part in parts)
        if positive <= amount_tolerance or negative_abs <= amount_tolerance:
            continue
        both_sign_keys += 1
        offset = min(positive, negative_abs)
        total_offset += offset
        for part in parts:
            row = part["warehouse_row"]
            result.append(
                {
                    "month": material["month"],
                    "code": material["code"],
                    "name": material["name"],
                    "spec": material["spec"],
                    "unit": material["unit"],
                    "pooled_avg": material["caats_avg"],
                    "material_issue_difference": material["issue_diff"],
                    "warehouse": row["warehouse"],
                    "warehouse_code": row["warehouse_code"],
                    "group": row["group"],
                    "warehouse_issue_quantity": row["oq"],
                    "warehouse_issue_amount": row["oa"],
                    "pooled_expected_amount": row["oq"] * material["caats_avg"],
                    "warehouse_contribution": part["contribution"],
                    "key_positive": positive,
                    "key_negative_abs": negative_abs,
                    "key_offset": offset,
                    "key_net": positive - negative_abs,
                }
            )
    result.sort(key=lambda row: (-row["key_offset"], row["month"], row["code"], -abs(row["warehouse_contribution"])))
    summary = {
        "multi_issue_keys": multi_issue_keys,
        "clean_multi_keys": clean_multi_keys,
        "both_sign_keys": both_sign_keys,
        "offset_rows": len(result),
        "total_offset": total_offset,
    }
    return result, summary


def prepare(args: argparse.Namespace) -> dict[str, Any]:
    project_root = args.project_root.resolve()
    staging_dir = args.staging_dir.resolve()
    staging_dir.mkdir(parents=True, exist_ok=True)
    converted_dir = staging_dir / "converted_xls"
    config = json.loads(args.config.read_text(encoding="utf-8"))
    periods = [str(value) for value in config["periods"]]
    amount_tolerance = float(config["amount_tolerance"])
    quantity_tolerance = float(config["quantity_tolerance"])
    special_document_types = set(config["special_document_types"])
    excluded_movement_categories = set(config.get("excluded_movement_categories", []))

    ledger_path = (args.ledger or project_root / "仓库管理-出入库流水账-2026.01-06.xlsx").resolve()
    summary_dir = (args.summary_dir or project_root / "收发存汇总").resolve()
    warehouse_master_source = args.warehouse_master.resolve()
    warehouse_master_path = convert_xls(warehouse_master_source, converted_dir / "warehouse_master", args.soffice)
    summary_sources = sorted(path for path in summary_dir.iterdir() if path.suffix.lower() in {".xls", ".xlsx"})
    converted_summaries = [convert_xls(path, converted_dir / "summaries", args.soffice) for path in summary_sources]

    master, master_by_name = read_master(warehouse_master_path)
    summaries, summary_by_key = read_summaries(converted_summaries, periods)

    workbook = load_workbook(ledger_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(min_row=2, values_only=True)
    headers = list(next(rows))
    ledger_index = required_columns(
        headers,
        ["日期", "单据类型", "单据号", "仓库", "存货编码", "存货名称", "规格型号", "收发类别", "记账人", "入库数量", "入库单价", "入库金额", "出库数量", "出库单价", "出库金额"],
        ledger_path.name,
    )

    material_aggregate: dict[tuple[str, str], defaultdict[str, float]] = defaultdict(lambda: defaultdict(float))
    excluded_material_aggregate: dict[tuple[str, str], defaultdict[str, float]] = defaultdict(lambda: defaultdict(float))
    excluded_category_aggregate: dict[str, defaultdict[str, float]] = defaultdict(lambda: defaultdict(float))
    warehouse_aggregate: dict[tuple[str, str, str], defaultdict[str, float]] = defaultdict(lambda: defaultdict(float))
    warehouse_categories: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    warehouse_document_types: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    group_aggregate: dict[tuple[str, str, str], defaultdict[str, float]] = defaultdict(lambda: defaultdict(float))
    group_categories: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    excluded_income_details: list[dict[str, Any]] = []
    excluded_issue_details: list[dict[str, Any]] = []
    valid_rows = posted_cost_rows = unposted_rows = no_cost_rows = 0
    missing_master: set[str] = set()

    for source_row, row in enumerate(rows, start=3):
        if row[ledger_index["存货编码"]] in (None, "") or row[ledger_index["日期"]] in (None, ""):
            continue
        valid_rows += 1
        month = text(row[ledger_index["日期"]])[:7].replace("-", "")
        if month not in periods:
            continue
        code = text(row[ledger_index["存货编码"]])
        warehouse = text(row[ledger_index["仓库"]])
        category = text(row[ledger_index["收发类别"]])
        document_type = text(row[ledger_index["单据类型"]])
        inbound_quantity = number(row[ledger_index["入库数量"]])
        inbound_amount = number(row[ledger_index["入库金额"]])
        outbound_quantity = number(row[ledger_index["出库数量"]])
        outbound_amount = number(row[ledger_index["出库金额"]])
        material_key = (month, code)
        if category in excluded_movement_categories:
            excluded = excluded_material_aggregate[material_key]
            excluded["rows"] += 1
            excluded["iq"] += inbound_quantity
            excluded["ia"] += inbound_amount
            excluded["oq"] += outbound_quantity
            excluded["oa"] += outbound_amount
            category_total = excluded_category_aggregate[category]
            category_total["rows"] += 1
            category_total["iq"] += inbound_quantity
            category_total["ia"] += inbound_amount
            category_total["oq"] += outbound_quantity
            category_total["oa"] += outbound_amount
        warehouse_master = master_by_name.get(warehouse)
        if warehouse_master is None:
            missing_master.add(warehouse)
            continue
        posted = row[ledger_index["记账人"]] not in (None, "")
        costed = warehouse_master.get("记入成本") == "是"
        warehouse_key = (month, warehouse, code)
        warehouse_categories[warehouse_key].add(category)
        warehouse_document_types[warehouse_key].add(document_type)
        warehouse_aggregate[warehouse_key]["all_rows"] += 1
        if category in excluded_movement_categories:
            for name, value in (("excluded_rows", 1), ("excluded_iq", inbound_quantity), ("excluded_ia", inbound_amount), ("excluded_oq", outbound_quantity), ("excluded_oa", outbound_amount)):
                warehouse_aggregate[warehouse_key][name] += value
        if not posted:
            unposted_rows += 1
            warehouse_aggregate[warehouse_key]["unposted_rows"] += 1
            warehouse_aggregate[warehouse_key]["unposted_iq"] += inbound_quantity
            warehouse_aggregate[warehouse_key]["unposted_oq"] += outbound_quantity
            continue
        if not costed:
            no_cost_rows += 1
            warehouse_aggregate[warehouse_key]["posted_no_cost_rows"] += 1
            warehouse_aggregate[warehouse_key]["posted_no_cost_iq"] += inbound_quantity
            warehouse_aggregate[warehouse_key]["posted_no_cost_oq"] += outbound_quantity
            continue
        posted_cost_rows += 1
        for aggregate, key in ((material_aggregate, material_key), (warehouse_aggregate, warehouse_key)):
            aggregate[key]["posted_rows"] += 1
            aggregate[key]["iq"] += inbound_quantity
            aggregate[key]["ia"] += inbound_amount
            aggregate[key]["oq"] += outbound_quantity
            aggregate[key]["oa"] += outbound_amount
        if category in excluded_movement_categories:
            for aggregate, key in ((material_aggregate, material_key), (warehouse_aggregate, warehouse_key)):
                for name, value in (("excluded_posted_rows", 1), ("excluded_posted_iq", inbound_quantity), ("excluded_posted_ia", inbound_amount), ("excluded_posted_oq", outbound_quantity), ("excluded_posted_oa", outbound_amount)):
                    aggregate[key][name] += value
            common_detail = {
                "month": month,
                "date": text(row[ledger_index["日期"]])[:10],
                "document_type": document_type,
                "document_number": text(row[ledger_index["单据号"]]),
                "warehouse_code": warehouse_master["仓库编码"],
                "warehouse": warehouse,
                "warehouse_group": text(warehouse_master.get("仓库核算组")),
                "code": code,
                "name": text(row[ledger_index.get("存货名称")]) if "存货名称" in ledger_index else "",
                "spec": text(row[ledger_index.get("规格型号")]) if "规格型号" in ledger_index else "",
                "category": category,
                "posted_by": text(row[ledger_index["记账人"]]),
                "source_file": ledger_path.name,
                "source_row": source_row,
            }
            if abs(inbound_quantity) > quantity_tolerance or abs(inbound_amount) > amount_tolerance:
                excluded_income_details.append(
                    {
                        **common_detail,
                        "quantity": inbound_quantity,
                        "unit_price": number(row[ledger_index.get("入库单价")]) if "入库单价" in ledger_index else 0.0,
                        "amount": inbound_amount,
                    }
                )
            if abs(outbound_quantity) > quantity_tolerance or abs(outbound_amount) > amount_tolerance:
                excluded_issue_details.append(
                    {
                        **common_detail,
                        "quantity": outbound_quantity,
                        "unit_price": number(row[ledger_index.get("出库单价")]) if "出库单价" in ledger_index else 0.0,
                        "amount": outbound_amount,
                    }
                )
        is_special = document_type in special_document_types
        if is_special:
            for name, value in (("special_rows", 1), ("special_iq", inbound_quantity), ("special_ia", inbound_amount), ("special_oq", outbound_quantity), ("special_oa", outbound_amount)):
                warehouse_aggregate[warehouse_key][name] += value
            group = text(warehouse_master.get("仓库核算组")) or "(无组)"
            group_key = (month, group, code)
            group_aggregate[group_key]["rows"] += 1
            group_aggregate[group_key]["iq"] += inbound_quantity
            group_aggregate[group_key]["ia"] += inbound_amount
            group_aggregate[group_key]["oq"] += outbound_quantity
            group_aggregate[group_key]["oa"] += outbound_amount
            group_categories[group_key].add(category)
        else:
            for name, value in (("normal_rows", 1), ("normal_iq", inbound_quantity), ("normal_ia", inbound_amount), ("normal_oq", outbound_quantity), ("normal_oa", outbound_amount)):
                warehouse_aggregate[warehouse_key][name] += value
        if category == "调拨入库":
            material_aggregate[material_key]["transfer_iq"] += inbound_quantity
            material_aggregate[material_key]["transfer_ia"] += inbound_amount
            warehouse_aggregate[warehouse_key]["transfer_iq"] += inbound_quantity
            warehouse_aggregate[warehouse_key]["transfer_ia"] += inbound_amount
        if category == "调拨出库":
            material_aggregate[material_key]["transfer_oq"] += outbound_quantity
            material_aggregate[material_key]["transfer_oa"] += outbound_amount
            warehouse_aggregate[warehouse_key]["transfer_oq"] += outbound_quantity
            warehouse_aggregate[warehouse_key]["transfer_oa"] += outbound_amount

    price_by_material: dict[tuple[str, str], list[float]] = defaultdict(list)
    warehouses_by_material: dict[tuple[str, str], set[str]] = defaultdict(set)
    issue_warehouses_by_material: dict[tuple[str, str], set[str]] = defaultdict(set)
    for (month, warehouse, code), aggregate in warehouse_aggregate.items():
        if aggregate.get("posted_rows", 0):
            warehouses_by_material[(month, code)].add(warehouse)
        filtered_oq = aggregate.get("oq", 0) - aggregate.get("excluded_posted_oq", 0)
        filtered_oa = aggregate.get("oa", 0) - aggregate.get("excluded_posted_oa", 0)
        if filtered_oq:
            issue_warehouses_by_material[(month, code)].add(warehouse)
            price_by_material[(month, code)].append(filtered_oa / filtered_oq)

    material_rows: list[dict[str, Any]] = []
    offset_material_rows: list[dict[str, Any]] = []
    for record in sorted(summaries, key=lambda item: (item["month"], item["code"])):
        key = (record["month"], record["code"])
        aggregate = material_aggregate.get(key, {})
        excluded = excluded_material_aggregate.get(key, {})
        excluded_posted = {
            "iq": aggregate.get("excluded_posted_iq", 0),
            "ia": aggregate.get("excluded_posted_ia", 0),
            "oq": aggregate.get("excluded_posted_oq", 0),
            "oa": aggregate.get("excluded_posted_oa", 0),
        }
        included = {
            "iq": aggregate.get("iq", 0) - aggregate.get("excluded_posted_iq", 0),
            "ia": aggregate.get("ia", 0) - aggregate.get("excluded_posted_ia", 0),
            "oq": aggregate.get("oq", 0) - aggregate.get("excluded_posted_oq", 0),
            "oa": aggregate.get("oa", 0) - aggregate.get("excluded_posted_oa", 0),
        }
        item = apply_movement_exclusions(record, included, excluded, excluded_posted)
        denominator = item["期初数量"] + included["iq"]
        average = (item["期初金额"] + included["ia"]) / denominator if abs(denominator) > 1e-12 else 0.0
        calculated_issue = included["oq"] * average
        calculated_end = item["期初金额"] + included["ia"] - calculated_issue
        issue_difference = item["u8_filtered_issue_amount"] - calculated_issue
        prices = price_by_material.get(key, [])
        minimum_price = min(prices) if prices else None
        maximum_price = max(prices) if prices else None
        income_gap = record["收入金额"] - aggregate.get("ia", 0)
        issue_gap = record["发出金额"] - aggregate.get("oa", 0)
        multi_price = len(issue_warehouses_by_material.get(key, set())) > 1 and minimum_price is not None and maximum_price is not None and abs(maximum_price - minimum_price) > amount_tolerance
        if multi_price:
            reason = "跨仓价格混合（已剔除四类移动）"
        elif within_amount_tolerance(issue_difference, amount_tolerance):
            reason = "U8剔除后与物料月CAATS重算一致"
        else:
            reason = "仓库/特殊单据/舍入或其他配置"
        item.update(
            {
                "ledger_iq": aggregate.get("iq", 0),
                "ledger_ia": aggregate.get("ia", 0),
                "ledger_oq": aggregate.get("oq", 0),
                "ledger_oa": aggregate.get("oa", 0),
                "filtered_ledger_iq": included["iq"],
                "filtered_ledger_ia": included["ia"],
                "filtered_ledger_oq": included["oq"],
                "filtered_ledger_oa": included["oa"],
                "transfer_iq": aggregate.get("transfer_iq", 0),
                "transfer_ia": aggregate.get("transfer_ia", 0),
                "transfer_oq": aggregate.get("transfer_oq", 0),
                "transfer_oa": aggregate.get("transfer_oa", 0),
                "warehouse_count": len(warehouses_by_material.get(key, set())),
                "issue_warehouse_count": len(issue_warehouses_by_material.get(key, set())),
                "warehouse_price_min": minimum_price,
                "warehouse_price_max": maximum_price,
                "warehouse_price_range": maximum_price - minimum_price if prices else None,
                "caats_avg": average,
                "caats_issue": calculated_issue,
                "issue_diff": issue_difference,
                "caats_end": calculated_end,
                "end_diff": item["u8_filtered_end_amount"] - calculated_end,
                "income_amount_gap": income_gap,
                "issue_amount_gap": issue_gap,
                "reason": reason,
                "status": "PASS" if within_amount_tolerance(issue_difference, amount_tolerance) else "REVIEW",
            }
        )
        material_rows.append(item)

        original_denominator = record["期初数量"] + record["收入数量"]
        original_average = (record["期初金额"] + record["收入金额"]) / original_denominator if abs(original_denominator) > 1e-12 else 0.0
        offset_material_rows.append(
            {
                **record,
                "caats_avg": original_average,
                "issue_diff": record["发出金额"] - record["发出数量"] * original_average,
                "issue_amount_gap": issue_gap,
            }
        )

    warehouse_rows: list[dict[str, Any]] = []
    for (month, warehouse, code), aggregate in sorted(warehouse_aggregate.items()):
        if not aggregate.get("all_rows", 0):
            continue
        warehouse_master = master_by_name[warehouse]
        summary = summary_by_key.get((month, code), {})
        warehouse_rows.append(
            {
                "month": month,
                "warehouse_code": warehouse_master["仓库编码"],
                "warehouse": warehouse,
                "group": text(warehouse_master.get("仓库核算组")),
                "method": text(warehouse_master.get("计价方式")),
                "costed": text(warehouse_master.get("记入成本")),
                "stopped": text(warehouse_master.get("停用日期")),
                "code": code,
                "name": summary.get("name", ""),
                "spec": summary.get("spec", ""),
                "unit": summary.get("unit", ""),
                "posted_rows": aggregate.get("posted_rows", 0),
                "unposted_rows": aggregate.get("unposted_rows", 0),
                "posted_no_cost_rows": aggregate.get("posted_no_cost_rows", 0),
                "iq": aggregate.get("iq", 0),
                "ia": aggregate.get("ia", 0),
                "oq": aggregate.get("oq", 0),
                "oa": aggregate.get("oa", 0),
                "out_unit": aggregate.get("oa", 0) / aggregate["oq"] if aggregate.get("oq", 0) else None,
                "normal_iq": aggregate.get("normal_iq", 0),
                "normal_ia": aggregate.get("normal_ia", 0),
                "normal_oq": aggregate.get("normal_oq", 0),
                "normal_oa": aggregate.get("normal_oa", 0),
                "special_iq": aggregate.get("special_iq", 0),
                "special_ia": aggregate.get("special_ia", 0),
                "special_oq": aggregate.get("special_oq", 0),
                "special_oa": aggregate.get("special_oa", 0),
                "transfer_iq": aggregate.get("transfer_iq", 0),
                "transfer_ia": aggregate.get("transfer_ia", 0),
                "transfer_oq": aggregate.get("transfer_oq", 0),
                "transfer_oa": aggregate.get("transfer_oa", 0),
                "excluded_rows": aggregate.get("excluded_rows", 0),
                "excluded_iq": aggregate.get("excluded_iq", 0),
                "excluded_ia": aggregate.get("excluded_ia", 0),
                "excluded_oq": aggregate.get("excluded_oq", 0),
                "excluded_oa": aggregate.get("excluded_oa", 0),
                "filtered_iq": aggregate.get("iq", 0) - aggregate.get("excluded_posted_iq", 0),
                "filtered_ia": aggregate.get("ia", 0) - aggregate.get("excluded_posted_ia", 0),
                "filtered_oq": aggregate.get("oq", 0) - aggregate.get("excluded_posted_oq", 0),
                "filtered_oa": aggregate.get("oa", 0) - aggregate.get("excluded_posted_oa", 0),
                "unposted_iq": aggregate.get("unposted_iq", 0),
                "unposted_oq": aggregate.get("unposted_oq", 0),
                "categories": "、".join(sorted(value for value in warehouse_categories[(month, warehouse, code)] if value)),
                "doc_types": "、".join(sorted(value for value in warehouse_document_types[(month, warehouse, code)] if value)),
                "u8_mat_iq": summary.get("收入数量", 0),
                "u8_mat_oq": summary.get("发出数量", 0),
                "mat_sum_iq": material_aggregate.get((month, code), {}).get("iq", 0),
                "mat_sum_oq": material_aggregate.get((month, code), {}).get("oq", 0),
            }
        )

    group_rows: list[dict[str, Any]] = []
    for (month, group, code), aggregate in sorted(group_aggregate.items()):
        summary = summary_by_key.get((month, code), {})
        group_rows.append(
            {
                "month": month,
                "group": group,
                "code": code,
                "name": summary.get("name", ""),
                "spec": summary.get("spec", ""),
                "unit": summary.get("unit", ""),
                "rows": aggregate.get("rows", 0),
                "iq": aggregate.get("iq", 0),
                "ia": aggregate.get("ia", 0),
                "oq": aggregate.get("oq", 0),
                "oa": aggregate.get("oa", 0),
                "in_unit": aggregate.get("ia", 0) / aggregate["iq"] if aggregate.get("iq", 0) else None,
                "out_unit": aggregate.get("oa", 0) / aggregate["oq"] if aggregate.get("oq", 0) else None,
                "categories": "、".join(sorted(value for value in group_categories[(month, group, code)] if value)),
            }
        )

    monthly: list[dict[str, Any]] = []
    for month in periods:
        rows_for_month = [row for row in material_rows if row["month"] == month]
        monthly.append(
            {
                "month": month,
                "rows": len(rows_for_month),
                "review": sum(row["status"] == "REVIEW" for row in rows_for_month),
                "pass": sum(row["status"] == "PASS" for row in rows_for_month),
                "signed_issue_diff": sum(row["issue_diff"] for row in rows_for_month),
                "abs_issue_diff": sum(abs(row["issue_diff"]) for row in rows_for_month),
            }
        )

    transition_checks = build_transition_checks(summaries, periods, quantity_tolerance, amount_tolerance)
    movement_by_key = {
        key: {
            "iq": aggregate.get("iq", 0) - aggregate.get("excluded_posted_iq", 0),
            "ia": aggregate.get("ia", 0) - aggregate.get("excluded_posted_ia", 0),
            "oq": aggregate.get("oq", 0) - aggregate.get("excluded_posted_oq", 0),
            "oa": aggregate.get("oa", 0) - aggregate.get("excluded_posted_oa", 0),
        }
        for key, aggregate in material_aggregate.items()
    }
    continuous_material_rows = build_complete_material_panel(
        material_rows, periods, movement_by_key, quantity_tolerance, amount_tolerance
    )
    continuous_monthly, continuous_rows = build_continuous_analysis(continuous_material_rows, periods, amount_tolerance)
    final_month = periods[-1]
    final_rows = [row for row in continuous_rows if row["month"] == final_month]
    offset_rows, offset_summary = build_offset_rows(offset_material_rows, warehouse_rows, quantity_tolerance, amount_tolerance)
    positive_difference = sum(max(row["issue_diff"], 0.0) for row in material_rows)
    negative_difference_abs = sum(max(-row["issue_diff"], 0.0) for row in material_rows)
    scope_issue_difference = sum(row["u8_filtered_issue_amount"] - row["filtered_ledger_oa"] for row in material_rows)
    repricing_issue_difference = sum(row["filtered_ledger_oa"] - row["caats_issue"] for row in material_rows)
    metrics = {
        "material_month_rows": len(material_rows),
        "warehouse_material_month_rows": len(warehouse_rows),
        "group_special_rows": len(group_rows),
        "valid_ledger_rows": valid_rows,
        "posted_cost_rows": posted_cost_rows,
        "unposted_rows": unposted_rows,
        "no_cost_rows": no_cost_rows,
        "missing_master": sorted(missing_master),
        "review_rows": sum(row["status"] == "REVIEW" for row in material_rows),
        "pass_rows": sum(row["status"] == "PASS" for row in material_rows),
        "signed_issue_diff": sum(row["issue_diff"] for row in material_rows),
        "abs_issue_diff": sum(abs(row["issue_diff"]) for row in material_rows),
        "max_abs_issue_diff": max(abs(row["issue_diff"]) for row in material_rows),
        "amount_gap_rows": sum(abs(row["income_amount_gap"]) > amount_tolerance or abs(row["issue_amount_gap"]) > amount_tolerance for row in material_rows),
        "continuous_panel_rows": len(continuous_material_rows),
        "continuous_synthetic_rows": sum(row.get("record_type") == "补齐无收发月份" for row in continuous_material_rows),
        "final_material_rows": len(final_rows),
        "final_synthetic_rows": sum(row["record_type"] == "补齐无收发月份" for row in final_rows),
        "final_u8_end_amount": sum(row["u8_filtered_end_amount"] for row in final_rows),
        "final_caats_end_amount": sum(row["continuous_caats_end"] for row in final_rows),
        "final_end_difference": sum(row["continuous_end_difference"] for row in final_rows),
        "rollforward_review_rows": sum(row["rollforward_status"] == "REVIEW" for row in continuous_rows),
        "excluded_movement_categories": sorted(excluded_movement_categories),
        "excluded_movement_rows": int(sum(value.get("rows", 0) for value in excluded_category_aggregate.values())),
        "excluded_movement_iq": sum(value.get("iq", 0) for value in excluded_category_aggregate.values()),
        "excluded_movement_ia": sum(value.get("ia", 0) for value in excluded_category_aggregate.values()),
        "excluded_movement_oq": sum(value.get("oq", 0) for value in excluded_category_aggregate.values()),
        "excluded_movement_oa": sum(value.get("oa", 0) for value in excluded_category_aggregate.values()),
        "excluded_income_detail_rows": len(excluded_income_details),
        "excluded_income_detail_quantity": sum(row["quantity"] for row in excluded_income_details),
        "excluded_income_detail_amount": sum(row["amount"] for row in excluded_income_details),
        "excluded_issue_detail_rows": len(excluded_issue_details),
        "excluded_issue_detail_quantity": sum(row["quantity"] for row in excluded_issue_details),
        "excluded_issue_detail_amount": sum(row["amount"] for row in excluded_issue_details),
        "positive_issue_diff": positive_difference,
        "negative_issue_diff_abs": negative_difference_abs,
        "global_netting": min(positive_difference, negative_difference_abs),
        "scope_issue_difference": scope_issue_difference,
        "repricing_issue_difference": repricing_issue_difference,
        "excluded_income_quantity_bridge": sum(row["收入数量"] - row["u8_filtered_income_quantity"] for row in material_rows),
        "excluded_income_amount_bridge": sum(row["收入金额"] - row["u8_filtered_income_amount"] for row in material_rows),
        "excluded_issue_quantity_bridge": sum(row["发出数量"] - row["u8_filtered_issue_quantity"] for row in material_rows),
        "excluded_issue_bridge": sum(row["发出金额"] - row["u8_filtered_issue_amount"] for row in material_rows),
    }
    input_files = [ledger_path, warehouse_master_source, *summary_sources, project_root / "CAATS交付模板.xlsx"]
    source_manifest = [
        {"role": "输入文件", "path": str(path), "name": path.name, "size": path.stat().st_size, "sha256": sha256(path)}
        for path in input_files
    ]
    return {
        "config": config,
        "metrics": metrics,
        "monthly": monthly,
        "continuous_monthly": continuous_monthly,
        "continuous_material_rows": continuous_material_rows,
        "continuous_difference_rows": sorted(
            (row for row in continuous_rows if row["status"] == "REVIEW"),
            key=lambda row: abs(row["continuous_issue_difference"]),
            reverse=True,
        ),
        "transition_checks": transition_checks,
        "offset_rows": offset_rows,
        "offset_summary": offset_summary,
        "material_rows": material_rows,
        "difference_rows": sorted(
            (row for row in material_rows if row["status"] == "REVIEW"),
            key=lambda row: abs(row["issue_diff"]),
            reverse=True,
        ),
        "warehouse_rows": warehouse_rows,
        "group_rows": group_rows,
        "master": master,
        "excluded_income_details": excluded_income_details,
        "excluded_issue_details": excluded_issue_details,
        "excluded_movement_breakdown": [
            {
                "category": category,
                "rows": int(excluded_category_aggregate[category].get("rows", 0)),
                "iq": excluded_category_aggregate[category].get("iq", 0),
                "ia": excluded_category_aggregate[category].get("ia", 0),
                "oq": excluded_category_aggregate[category].get("oq", 0),
                "oa": excluded_category_aggregate[category].get("oa", 0),
            }
            for category in sorted(excluded_movement_categories)
        ],
        "source_manifest": source_manifest,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--warehouse-master", type=Path, required=True)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--soffice", type=Path, required=True)
    parser.add_argument("--staging-dir", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    parser.add_argument("--ledger", type=Path)
    parser.add_argument("--summary-dir", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    payload = prepare(args)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":"), default=str), encoding="utf-8")
    print(json.dumps({"metrics": payload["metrics"], "offset_summary": payload["offset_summary"], "output_json": str(args.output_json)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
