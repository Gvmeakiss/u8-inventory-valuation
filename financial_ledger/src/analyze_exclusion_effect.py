#!/usr/bin/env python3
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def n(value):
    return float(value or 0.0)


def run(rows, income_qty, income_amt, issue_qty, issue_amt, end_amt):
    by_code = defaultdict(list)
    for row in rows:
        by_code[str(row["code"])].append(row)
    details = []
    monthly = defaultdict(lambda: defaultdict(float))
    for code, code_rows in by_code.items():
        prior = None
        cumulative = defaultdict(float)
        for row in sorted(code_rows, key=lambda item: str(item["month"])):
            begin_qty = prior["end_qty"] if prior else n(row["期初数量"])
            begin_amt = prior["end_amt"] if prior else n(row["期初金额"])
            iq = n(row[income_qty])
            ia = n(row[income_amt])
            oq = n(row[issue_qty])
            denominator = begin_qty + iq
            average = (begin_amt + ia) / denominator if abs(denominator) > 1e-12 else 0.0
            caats_issue = oq * average
            caats_end_qty = begin_qty + iq - oq
            caats_end_amt = begin_amt + ia - caats_issue
            issue_diff = n(row[issue_amt]) - caats_issue
            end_diff = n(row[end_amt]) - caats_end_amt
            for field in ("excluded_posted_iq", "excluded_posted_ia", "excluded_posted_oq", "excluded_posted_oa"):
                cumulative[field] += n(row.get(field))
            detail = {
                "month": str(row["month"]),
                "code": code,
                "name": row.get("name", ""),
                "begin_qty": begin_qty,
                "begin_amt": begin_amt,
                "average": average,
                "caats_issue": caats_issue,
                "issue_diff": issue_diff,
                "caats_end_qty": caats_end_qty,
                "caats_end_amt": caats_end_amt,
                "end_diff": end_diff,
                "current_excluded_iq": n(row.get("excluded_posted_iq")),
                "current_excluded_ia": n(row.get("excluded_posted_ia")),
                "current_excluded_oq": n(row.get("excluded_posted_oq")),
                "current_excluded_oa": n(row.get("excluded_posted_oa")),
                **{f"cum_{key}": value for key, value in cumulative.items()},
            }
            details.append(detail)
            month = monthly[detail["month"]]
            month["issue_diff"] += issue_diff
            month["issue_abs"] += abs(issue_diff)
            month["end_diff"] += end_diff
            month["end_abs"] += abs(end_diff)
            prior = {"end_qty": caats_end_qty, "end_amt": caats_end_amt}
    return details, monthly


def independent(rows, income_qty, income_amt, issue_qty, issue_amt):
    signed = absolute = 0.0
    review = 0
    details = []
    for row in rows:
        denominator = n(row["期初数量"]) + n(row[income_qty])
        average = (n(row["期初金额"]) + n(row[income_amt])) / denominator if abs(denominator) > 1e-12 else 0.0
        caats_issue = n(row[issue_qty]) * average
        difference = n(row[issue_amt]) - caats_issue
        signed += difference
        absolute += abs(difference)
        review += abs(round(difference, 10)) > 0.01
        details.append((str(row["month"]), str(row["code"]), difference))
    return {"signed": signed, "absolute": absolute, "review": review, "pass": len(rows) - review}, details


def main():
    data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    rows = data["material_rows"]
    old_independent, _ = independent(rows, "ledger_iq", "ledger_ia", "ledger_oq", "发出金额")
    new_independent, _ = independent(rows, "filtered_ledger_iq", "filtered_ledger_ia", "filtered_ledger_oq", "u8_filtered_issue_amount")

    old_details, old_monthly = run(rows, "ledger_iq", "ledger_ia", "ledger_oq", "发出金额", "结存金额")
    new_details, new_monthly = run(rows, "filtered_ledger_iq", "filtered_ledger_ia", "filtered_ledger_oq", "u8_filtered_issue_amount", "u8_filtered_end_amount")
    old_by_key = {(row["month"], row["code"]): row for row in old_details}
    new_by_key = {(row["month"], row["code"]): row for row in new_details}
    comparisons = []
    for key in sorted(old_by_key.keys() & new_by_key.keys()):
        old = old_by_key[key]
        new = new_by_key[key]
        if key[0] != "202606":
            continue
        comparisons.append({
            "month": key[0], "code": key[1], "name": new["name"],
            "old_end_diff": old["end_diff"], "new_end_diff": new["end_diff"],
            "change": new["end_diff"] - old["end_diff"],
            "old_average": old["average"], "new_average": new["average"],
            "old_caats_end": old["caats_end_amt"], "new_caats_end": new["caats_end_amt"],
            "u8_benchmark_change": (
                n(next(row for row in rows if str(row["month"]) == key[0] and str(row["code"]) == key[1])["u8_filtered_end_amount"])
                - n(next(row for row in rows if str(row["month"]) == key[0] and str(row["code"]) == key[1])["结存金额"])
            ),
            "cum_excluded_iq": new["cum_excluded_posted_iq"],
            "cum_excluded_ia": new["cum_excluded_posted_ia"],
            "cum_excluded_oq": new["cum_excluded_posted_oq"],
            "cum_excluded_oa": new["cum_excluded_posted_oa"],
            "current_excluded_ia": new["current_excluded_ia"],
            "current_excluded_oa": new["current_excluded_oa"],
        })

    exclusion_monthly = []
    for month in sorted({str(row["month"]) for row in rows}):
        month_rows = [row for row in rows if str(row["month"]) == month]
        exclusion_monthly.append({
            "month": month,
            "iq": sum(n(row.get("excluded_posted_iq")) for row in month_rows),
            "ia": sum(n(row.get("excluded_posted_ia")) for row in month_rows),
            "oq": sum(n(row.get("excluded_posted_oq")) for row in month_rows),
            "oa": sum(n(row.get("excluded_posted_oa")) for row in month_rows),
            "amount_net_in_minus_out": sum(n(row.get("excluded_posted_ia")) - n(row.get("excluded_posted_oa")) for row in month_rows),
            "quantity_net_in_minus_out": sum(n(row.get("excluded_posted_iq")) - n(row.get("excluded_posted_oq")) for row in month_rows),
            "material_months_not_closed_amount": sum(abs(n(row.get("excluded_posted_ia")) - n(row.get("excluded_posted_oa"))) > 0.01 for row in month_rows),
            "material_months_not_closed_quantity": sum(abs(n(row.get("excluded_posted_iq")) - n(row.get("excluded_posted_oq"))) > 0.000001 for row in month_rows),
        })

    june_old = old_monthly["202606"]
    june_new = new_monthly["202606"]
    result = {
        "independent": {"old": old_independent, "new": new_independent},
        "continuous_june": {
            "old": dict(june_old), "new": dict(june_new),
            "net_change": june_new["end_diff"] - june_old["end_diff"],
            "abs_change": june_new["end_abs"] - june_old["end_abs"],
        },
        "excluded_by_month": exclusion_monthly,
        "june_benchmark_change_total": sum(row["u8_benchmark_change"] for row in comparisons),
        "june_caats_end_change_total": sum(row["new_caats_end"] - row["old_caats_end"] for row in comparisons),
        "top_june_changes": sorted(comparisons, key=lambda row: abs(row["change"]), reverse=True)[:25],
        "june_change_signs": {
            "increase_abs_rows": sum(abs(row["new_end_diff"]) > abs(row["old_end_diff"]) + 0.01 for row in comparisons),
            "decrease_abs_rows": sum(abs(row["new_end_diff"]) + 0.01 < abs(row["old_end_diff"]) for row in comparisons),
            "unchanged_rows": sum(abs(abs(row["new_end_diff"]) - abs(row["old_end_diff"])) <= 0.01 for row in comparisons),
            "rows": len(comparisons),
        },
    }
    if len(sys.argv) > 3:
        old_path = Path(sys.argv[3])
        old_formula = load_workbook(old_path, read_only=True, data_only=False)
        old_values = load_workbook(old_path, read_only=True, data_only=True)
        result["old_workbook"] = {"path": str(old_path), "sheets": {}}
        for sheet_name in old_formula.sheetnames:
            ws_formula = old_formula[sheet_name]
            ws_values = old_values[sheet_name]
            ws_formula.reset_dimensions()
            ws_values.reset_dimensions()
            formula_rows = ws_formula.iter_rows(values_only=True)
            value_rows = ws_values.iter_rows(values_only=True)
            header_row = next(formula_rows, ())
            row2_formula = next(formula_rows, ())
            top_values = []
            for index, value_row in enumerate(value_rows):
                if index < 20:
                    top_values.append(list(value_row[:15]))
                if index >= 19:
                    break
            max_column = len(header_row)
            headers = list(header_row)
            result["old_workbook"]["sheets"][sheet_name] = {
                "rows": None,
                "columns": max_column,
                "headers": headers,
                "row2_formulas": list(row2_formula),
                "top_values": top_values,
            }
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(payload, encoding="utf-8")
    else:
        print(payload)


if __name__ == "__main__":
    main()
